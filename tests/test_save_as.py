"""Integration tests for Save As tracking."""
import os
import shutil
from pathlib import Path

import pytest
from docx import Document

from dochistory.core.database import Database
from dochistory.core.version_manager import VersionManager


@pytest.fixture
def setup(tmp_path):
    """Create DB, monitor dir, source file with multiple versions."""
    db = Database(str(tmp_path / "test.db"))
    manager = VersionManager(db)
    monitor_dir = tmp_path / "monitored"
    monitor_dir.mkdir()
    folder_id = db.add_folder(str(monitor_dir))

    source_path = monitor_dir / "source.docx"
    doc = Document()
    doc.add_paragraph("Version 1 content")
    doc.save(str(source_path))
    manager.create_version(str(source_path), folder_id)

    doc = Document(str(source_path))
    doc.add_paragraph("Version 2 content")
    doc.save(str(source_path))
    manager.create_version(str(source_path), folder_id)

    return manager, db, monitor_dir, folder_id, source_path


class TestSaveAsDetection:
    def test_find_source_by_hash(self, setup):
        manager, db, monitor_dir, folder_id, source_path = setup
        # Copy source to new file (simulates Save As)
        new_path = monitor_dir / "copy.docx"
        shutil.copyfile(str(source_path), str(new_path))

        source = manager.find_save_as_source(str(new_path))
        assert source is not None
        assert source["relative_path"] == "source.docx"

    def test_no_source_for_unique_file(self, setup, tmp_path):
        manager, db, monitor_dir, folder_id, source_path = setup
        new_path = monitor_dir / "unique.docx"
        doc = Document()
        doc.add_paragraph("Completely unique content")
        doc.save(str(new_path))

        source = manager.find_save_as_source(str(new_path))
        assert source is None


class TestAssociateHistory:
    def test_associate_copies_versions(self, setup):
        manager, db, monitor_dir, folder_id, source_path = setup
        # Create new file (Save As)
        new_path = monitor_dir / "copy.docx"
        shutil.copyfile(str(source_path), str(new_path))
        manager.handle_save_as(str(new_path), folder_id)

        new_file = db.find_file(folder_id, "copy.docx")
        source_file = db.find_file(folder_id, "source.docx")
        source_versions = db.get_versions(source_file["id"])

        manager.associate_history(new_file["id"], source_file["id"])

        new_versions = db.get_versions(new_file["id"])
        # Should have: 1 original + 2 inherited = 3
        assert len(new_versions) >= 3

    def test_source_file_unchanged_after_associate(self, setup):
        manager, db, monitor_dir, folder_id, source_path = setup
        new_path = monitor_dir / "copy.docx"
        shutil.copyfile(str(source_path), str(new_path))
        manager.handle_save_as(str(new_path), folder_id)

        source_file = db.find_file(folder_id, "source.docx")
        source_versions_before = len(db.get_versions(source_file["id"]))

        new_file = db.find_file(folder_id, "copy.docx")
        manager.associate_history(new_file["id"], source_file["id"])

        source_versions_after = len(db.get_versions(source_file["id"]))
        assert source_versions_after == source_versions_before

    def test_independent_evolution_after_associate(self, setup):
        manager, db, monitor_dir, folder_id, source_path = setup
        new_path = monitor_dir / "copy.docx"
        shutil.copyfile(str(source_path), str(new_path))
        manager.handle_save_as(str(new_path), folder_id)

        new_file = db.find_file(folder_id, "copy.docx")
        source_file = db.find_file(folder_id, "source.docx")
        manager.associate_history(new_file["id"], source_file["id"])

        # Modify source file
        doc = Document(str(source_path))
        doc.add_paragraph("Source only change")
        doc.save(str(source_path))
        manager.create_version(str(source_path), folder_id)

        # Modify new file
        doc = Document(str(new_path))
        doc.add_paragraph("Copy only change")
        doc.save(str(new_path))
        manager.create_version(str(new_path), folder_id)

        source_versions = db.get_versions(source_file["id"])
        new_versions = db.get_versions(new_file["id"])
        # Source should have 3 (2 original + 1 new), new should have 4 (1 + 2 inherited + 1 new)
        assert len(source_versions) == 3
        assert len(new_versions) >= 4
