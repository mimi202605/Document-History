"""Integration tests for the complete rollback flow."""
import os
from pathlib import Path

import pytest
from docx import Document

from dochistory.core.database import Database
from dochistory.core.version_manager import VersionManager


@pytest.fixture
def setup(tmp_path):
    db = Database(str(tmp_path / "test.db"))
    manager = VersionManager(db)
    monitor_dir = tmp_path / "monitored"
    monitor_dir.mkdir()
    folder_id = db.add_folder(str(monitor_dir))

    docx_path = monitor_dir / "report.docx"
    doc = Document()
    doc.add_paragraph("Version 1")
    doc.save(str(docx_path))
    v1 = manager.create_version(str(docx_path), folder_id)

    doc = Document(str(docx_path))
    doc.add_paragraph("Version 2")
    doc.save(str(docx_path))
    v2 = manager.create_version(str(docx_path), folder_id)

    doc = Document(str(docx_path))
    doc.add_paragraph("Version 3")
    doc.save(str(docx_path))
    v3 = manager.create_version(str(docx_path), folder_id)

    return manager, db, monitor_dir, folder_id, docx_path, [v1, v2, v3]


class TestRollbackFlow:
    def test_rollback_to_v1(self, setup):
        manager, db, monitor_dir, folder_id, docx_path, versions = setup
        v1, v2, v3 = versions
        file_rec = db.find_file(folder_id, "report.docx")

        manager.rollback_to_version(file_rec["id"], v1["id"])

        doc = Document(str(docx_path))
        texts = [p.text for p in doc.paragraphs]
        assert "Version 1" in texts
        assert "Version 2" not in texts
        assert "Version 3" not in texts

    def test_rollback_creates_backup(self, setup):
        manager, db, monitor_dir, folder_id, docx_path, versions = setup
        v1, v2, v3 = versions
        file_rec = db.find_file(folder_id, "report.docx")
        versions_before = len(db.get_versions(file_rec["id"]))

        manager.rollback_to_version(file_rec["id"], v1["id"])

        versions_after = len(db.get_versions(file_rec["id"]))
        assert versions_after == versions_before + 1

    def test_rollback_then_continue_editing(self, setup):
        manager, db, monitor_dir, folder_id, docx_path, versions = setup
        v1, v2, v3 = versions
        file_rec = db.find_file(folder_id, "report.docx")

        manager.rollback_to_version(file_rec["id"], v1["id"])

        # Continue editing after rollback
        doc = Document(str(docx_path))
        doc.add_paragraph("After rollback")
        doc.save(str(docx_path))
        v_new = manager.create_version(str(docx_path), folder_id)

        assert v_new is not None
        assert v_new["version_number"] > v3["version_number"]

    def test_export_version(self, setup, tmp_path):
        manager, db, monitor_dir, folder_id, docx_path, versions = setup
        v1, v2, v3 = versions
        file_rec = db.find_file(folder_id, "report.docx")

        export_path = tmp_path / "exported.docx"
        manager.export_version(file_rec["id"], v1["id"], str(export_path))

        assert export_path.exists()
        doc = Document(str(export_path))
        texts = [p.text for p in doc.paragraphs]
        assert "Version 1" in texts
        assert "Version 2" not in texts


class TestRollbackAtomicity:
    """缺陷3: rollback_to_version 在原子写入前创建备份，写入失败导致备份残留。"""

    def test_no_backup_residue_when_write_fails(self, setup, monkeypatch):
        manager, db, monitor_dir, folder_id, docx_path, versions = setup
        v1, v2, v3 = versions
        file_rec = db.find_file(folder_id, "report.docx")
        file_id = file_rec["id"]

        versions_before = len(db.get_versions(file_id))

        # mock _atomic_write 抛出 PermissionError 模拟写入失败
        monkeypatch.setattr(
            manager,
            "_atomic_write",
            lambda *args: (_ for _ in ()).throw(PermissionError("locked")),
        )

        # 调用 rollback_to_version 应抛出 PermissionError（不吞掉异常）
        with pytest.raises(PermissionError):
            manager.rollback_to_version(file_id, v1["id"])

        # 验证版本数量未增加（无备份残留）
        versions_after = len(db.get_versions(file_id))
        assert versions_after == versions_before

        # 取消 mock，验证 rollback_to_version 正常工作且创建备份
        monkeypatch.undo()
        manager.rollback_to_version(file_id, v1["id"])
        versions_after_ok = len(db.get_versions(file_id))
        assert versions_after_ok == versions_before + 1
