"""Tests for the VersionManager."""
import os
from pathlib import Path

import pytest
from docx import Document

from dochistory.core.database import Database
from dochistory.core.version_manager import VersionManager


@pytest.fixture
def vm(tmp_path):
    """Create a VersionManager with a temp database."""
    db_path = tmp_path / "test.db"
    db = Database(str(db_path))
    manager = VersionManager(db)
    yield manager, db
    db.close()


@pytest.fixture
def monitored_file(vm, tmp_path):
    """Create a monitored folder with a docx file."""
    manager, db = vm
    monitor_dir = tmp_path / "monitored"
    monitor_dir.mkdir()
    folder_id = db.add_folder(str(monitor_dir))

    docx_path = monitor_dir / "report.docx"
    doc = Document()
    doc.add_paragraph("Initial content")
    doc.save(str(docx_path))

    return monitor_dir, folder_id, docx_path


class TestCreateVersion:
    def test_create_first_version(self, vm, monitored_file):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        version = manager.create_version(str(docx_path), folder_id)
        assert version is not None
        assert version["version_number"] == 1

    def test_create_second_version(self, vm, monitored_file):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        manager.create_version(str(docx_path), folder_id)
        # Modify file
        doc = Document(str(docx_path))
        doc.add_paragraph("New content")
        doc.save(str(docx_path))
        version = manager.create_version(str(docx_path), folder_id)
        assert version["version_number"] == 2

    def test_skip_unchanged_file(self, vm, monitored_file):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        v1 = manager.create_version(str(docx_path), folder_id)
        v2 = manager.create_version(str(docx_path), folder_id)
        assert v2 is None  # No new version for unchanged file

    def test_manual_version_with_note(self, vm, monitored_file):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        version = manager.create_version(
            str(docx_path), folder_id, note="Manual checkpoint", is_auto=False
        )
        assert version["note"] == "Manual checkpoint"
        assert version["is_auto"] == 0


class TestGetVersions:
    def test_get_versions_ordered_newest_first(self, vm, monitored_file):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        manager.create_version(str(docx_path), folder_id)
        doc = Document(str(docx_path))
        doc.add_paragraph("v2")
        doc.save(str(docx_path))
        manager.create_version(str(docx_path), folder_id)

        file_rec = db.find_file(folder_id, "report.docx")
        versions = manager.get_versions(file_rec["id"])
        assert len(versions) == 2
        assert versions[0]["version_number"] == 2


class TestRollback:
    def test_rollback_restores_content(self, vm, monitored_file):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        v1 = manager.create_version(str(docx_path), folder_id)

        # Modify and create v2
        doc = Document(str(docx_path))
        doc.add_paragraph("Modified content")
        doc.save(str(docx_path))
        manager.create_version(str(docx_path), folder_id)

        # Rollback to v1
        file_rec = db.find_file(folder_id, "report.docx")
        manager.rollback_to_version(file_rec["id"], v1["id"])

        # Verify content restored
        restored_doc = Document(str(docx_path))
        texts = [p.text for p in restored_doc.paragraphs]
        assert "Initial content" in texts
        assert "Modified content" not in texts

    def test_rollback_creates_backup(self, vm, monitored_file):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        v1 = manager.create_version(str(docx_path), folder_id)

        doc = Document(str(docx_path))
        doc.add_paragraph("Modified")
        doc.save(str(docx_path))
        manager.create_version(str(docx_path), folder_id)

        file_rec = db.find_file(folder_id, "report.docx")
        versions_before = len(manager.get_versions(file_rec["id"]))
        manager.rollback_to_version(file_rec["id"], v1["id"])
        versions_after = len(manager.get_versions(file_rec["id"]))
        # Should have backup + original 2 versions
        assert versions_after == versions_before + 1


class TestExportVersion:
    def test_export_creates_file(self, vm, monitored_file, tmp_path):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        v1 = manager.create_version(str(docx_path), folder_id)

        export_path = tmp_path / "exported.docx"
        file_rec = db.find_file(folder_id, "report.docx")
        manager.export_version(file_rec["id"], v1["id"], str(export_path))
        assert export_path.exists()

        # Verify exported content
        doc = Document(str(export_path))
        texts = [p.text for p in doc.paragraphs]
        assert "Initial content" in texts


class TestCleanup:
    def test_cleanup_removes_old_versions(self, vm, monitored_file):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file
        for i in range(5):
            doc = Document(str(docx_path))
            doc.add_paragraph(f"version {i}")
            doc.save(str(docx_path))
            manager.create_version(str(docx_path), folder_id)

        file_rec = db.find_file(folder_id, "report.docx")
        manager.cleanup_old_versions(file_rec["id"], max_versions=3)
        versions = manager.get_versions(file_rec["id"])
        assert len(versions) == 3


class TestCreateVersionReadFailure:
    """缺陷1: create_version 在读取快照前提交哈希，读取失败导致数据丢失。"""

    def test_hash_not_updated_when_read_fails(self, vm, monitored_file, monkeypatch):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file

        # 创建初始版本 v1
        manager.create_version(str(docx_path), folder_id)
        file_rec = db.find_file(folder_id, "report.docx")
        old_hash = file_rec["file_hash"]

        # 修改文件内容（产生新哈希）
        doc = Document(str(docx_path))
        doc.add_paragraph("Modified content")
        doc.save(str(docx_path))

        # mock _read_with_retry 返回 None 模拟读取失败
        monkeypatch.setattr(manager, "_read_with_retry", lambda *args, **kwargs: None)

        # 调用 create_version 应返回 None
        result = manager.create_version(str(docx_path), folder_id)
        assert result is None

        # 验证数据库中 file_hash 仍为旧哈希（未被提前提交）
        file_rec_after = db.find_file(folder_id, "report.docx")
        assert file_rec_after["file_hash"] == old_hash

        # 取消 mock，再次调用 create_version 应成功创建版本
        # 证明哈希未被提前提交，状态未被永久跳过
        monkeypatch.undo()
        result = manager.create_version(str(docx_path), folder_id)
        assert result is not None
        assert result["version_number"] == 2


class TestRenameVersionUnreadable:
    """缺陷2: _record_rename_version 存储 b""（未压缩），回滚/导出时崩溃。"""

    def test_rename_version_with_unreadable_file_rollback_succeeds(
        self, vm, monitored_file, tmp_path, monkeypatch
    ):
        manager, db = vm
        monitor_dir, folder_id, docx_path = monitored_file

        # 创建初始版本
        manager.create_version(str(docx_path), folder_id)
        file_rec = db.find_file(folder_id, "report.docx")
        file_id = file_rec["id"]

        # 在磁盘上重命名文件，使 handle_rename 后新路径仍存在
        new_path = monitor_dir / "report_renamed.docx"
        os.rename(str(docx_path), str(new_path))

        # mock _read_with_retry 返回 None 模拟文件不可读
        monkeypatch.setattr(manager, "_read_with_retry", lambda *args, **kwargs: None)

        # 调用 handle_rename 触发 _record_rename_version
        manager.handle_rename(str(docx_path), str(new_path), folder_id)

        # 获取版本列表，找到 note 包含"重命名"的版本
        versions = db.get_versions(file_id)
        rename_version = None
        for v in versions:
            if v["note"] and "重命名" in v["note"]:
                rename_version = v
                break
        assert rename_version is not None

        # 验证 export_version 不崩溃，导出文件内容为空
        export_path = tmp_path / "exported_rename.docx"
        result = manager.export_version(file_id, rename_version["id"], str(export_path))
        assert result is True
        assert export_path.exists()
        assert export_path.read_bytes() == b""

        # 取消 mock，验证 rollback_to_version 到该重命名版本不崩溃
        monkeypatch.undo()
        manager.rollback_to_version(file_id, rename_version["id"])
