"""Integration tests for FileMonitor."""
import time
from pathlib import Path

import pytest
from docx import Document

from dochistory.core.file_monitor import FileMonitor


# PyQt6 signals require a QApplication instance
@pytest.fixture(scope="session", autouse=True)
def qapp():
    from PyQt6.QtWidgets import QApplication
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    yield app


def _wait(seconds: float):
    """Sleep while pumping the Qt event loop so cross-thread signals are delivered."""
    from PyQt6.QtWidgets import QApplication
    app = QApplication.instance()
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        if app is not None:
            app.processEvents()
        time.sleep(0.02)


@pytest.fixture
def monitor():
    """Create a FileMonitor instance."""
    m = FileMonitor()
    yield m
    m.stop_all()


class TestFileMonitor:
    def test_detect_modified_docx(self, monitor, temp_monitor_dir):
        events = []
        monitor.file_saved.connect(lambda p: events.append(p))

        monitor.add_watch(str(temp_monitor_dir))
        docx_path = temp_monitor_dir / "test.docx"
        Document().save(str(docx_path))

        # Wait for debounce + processing
        _wait(1.5)
        assert any("test.docx" in e for e in events)

    def test_ignore_temp_files(self, monitor, temp_monitor_dir):
        events = []
        monitor.file_saved.connect(lambda p: events.append(p))

        monitor.add_watch(str(temp_monitor_dir))
        # Create a temp file
        (temp_monitor_dir / "~$test.docx").write_text("temp")

        _wait(1.5)
        assert len(events) == 0

    def test_ignore_unsupported_files(self, monitor, temp_monitor_dir):
        events = []
        monitor.file_saved.connect(lambda p: events.append(p))

        monitor.add_watch(str(temp_monitor_dir))
        (temp_monitor_dir / "notes.txt").write_text("text")

        _wait(1.5)
        assert len(events) == 0

    def test_debounce_multiple_saves(self, monitor, temp_monitor_dir):
        events = []
        monitor.file_saved.connect(lambda p: events.append(p))

        monitor.add_watch(str(temp_monitor_dir))
        docx_path = temp_monitor_dir / "test.docx"
        doc = Document()
        doc.save(str(docx_path))

        # Rapidly modify
        for i in range(3):
            doc = Document()
            doc.add_paragraph(f"line {i}")
            doc.save(str(docx_path))
            _wait(0.1)

        _wait(1.5)
        # Should only have 1-2 events due to debounce (initial + final)
        assert len(events) <= 2

    def test_detect_rename(self, monitor, temp_monitor_dir):
        moved_events = []
        monitor.file_moved.connect(lambda s, d: moved_events.append((s, d)))

        monitor.add_watch(str(temp_monitor_dir))
        docx_path = temp_monitor_dir / "old.docx"
        Document().save(str(docx_path))
        _wait(0.5)

        new_path = temp_monitor_dir / "new.docx"
        docx_path.rename(new_path)

        _wait(1.5)
        assert len(moved_events) == 1
        assert "old.docx" in moved_events[0][0]
        assert "new.docx" in moved_events[0][1]

    def test_detect_new_file(self, monitor, temp_monitor_dir):
        created_events = []
        monitor.file_created.connect(lambda p: created_events.append(p))

        monitor.add_watch(str(temp_monitor_dir))
        docx_path = temp_monitor_dir / "new.docx"
        Document().save(str(docx_path))

        _wait(1.5)
        assert any("new.docx" in e for e in created_events)

    def test_pause_resume(self, monitor, temp_monitor_dir):
        events = []
        monitor.file_saved.connect(lambda p: events.append(p))

        monitor.add_watch(str(temp_monitor_dir))
        monitor.pause()

        docx_path = temp_monitor_dir / "test.docx"
        Document().save(str(docx_path))

        _wait(1.5)
        assert len(events) == 0

        monitor.resume()
        doc = Document(str(docx_path))
        doc.add_paragraph("new")
        doc.save(str(docx_path))

        _wait(1.5)
        assert len(events) > 0

    def test_remove_watch(self, monitor, temp_monitor_dir):
        events = []
        monitor.file_saved.connect(lambda p: events.append(p))

        monitor.add_watch(str(temp_monitor_dir))
        monitor.remove_watch(str(temp_monitor_dir))

        docx_path = temp_monitor_dir / "test.docx"
        Document().save(str(docx_path))

        _wait(1.5)
        assert len(events) == 0
